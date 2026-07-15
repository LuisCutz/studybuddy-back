import io
import fitz
from docx import Document as DocxDocument
from pptx import Presentation

async def extract_text_from_bytes(file_bytes: bytes, filename: str) -> str:
    # Extrae texto de PDFs, Word, PowerPoint y archivos de texto plano.
    ext = filename.split('.')[-1].lower()
    text = ""
    
    try:
        if ext == "pdf":
            with fitz.open(stream=file_bytes, filetype="pdf") as doc:
                for page in doc:
                    text += page.get_text()
                    
        elif ext == "docx":
            doc = DocxDocument(io.BytesIO(file_bytes))
            text = "\n".join([para.text for para in doc.paragraphs])
            
        elif ext == "pptx":
            prs = Presentation(io.BytesIO(file_bytes))
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
                        
        elif ext == "txt":
            text = file_bytes.decode('utf-8', errors='ignore')
            
        else:
            raise ValueError(f"Formato no soportado para extracción: {ext}")
            
    except Exception as e:
        raise RuntimeError(f"Error procesando el archivo {filename}: {str(e)}")
        
    return text